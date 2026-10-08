import importlib.util
import json
import tempfile
import unittest
from pathlib import Path

from rag_kaggle.config import PipelineConfig
from rag_kaggle.models import Block, ParsedDocument
from rag_kaggle.profile import attach_profile_to_chunks, build_document_profile, normalize_key
from rag_kaggle.chunking import StructureAwareChunker
from rag_kaggle.storage import MetadataStore, bm25_text, tokenize_vi


def has(*modules):
    return all(importlib.util.find_spec(module) is not None for module in modules)


def block(block_id, block_type, content, order, **kwargs):
    return Block(block_id, "doc1", block_type, content, "qd.pdf", reading_order=order, **kwargs)


def official_document():
    document = ParsedDocument("doc1", "qd.pdf", "pdf", "hash")
    document.blocks = [
        block("b0", "text", "ỦY BAN NHÂN DÂN TỈNH HÒA BÌNH\nSố: 123/QĐ-UBND\nHòa Bình, ngày 5 tháng 3 năm 2025", 0, page=1),
        block("b1", "heading", "QUYẾT ĐỊNH về việc phê duyệt dự án", 1, page=1, section_path=["QUYẾT ĐỊNH"]),
        block(
            "b2", "text",
            "Căn cứ đề nghị của ông Nguyễn Văn An, Chủ tịch xã Hòa Bình. Dự án cầu Hòa Bình được phê duyệt. "
            "Nguyễn Văn An chịu trách nhiệm triển khai tại xã Hòa Bình.",
            2, page=1, section_path=["QUYẾT ĐỊNH"],
        ),
        block("b3", "text", "TM. ỦY BAN NHÂN DÂN\nCHỦ TỊCH\nTrần Thị Mai", 3, page=2),
    ]
    return document


class ProfileTests(unittest.TestCase):
    def test_profile_extracts_metadata_signers_and_entities(self):
        profile = build_document_profile(official_document())
        self.assertEqual("123/QĐ-UBND", profile["doc_number"])
        self.assertEqual("2025-03-05", profile["issue_date"])
        self.assertIn("ỦY BAN NHÂN DÂN", profile["issuer"])
        self.assertIn("Trần Thị Mai", profile["signers"])
        by_key = {entity["key"]: entity for entity in profile["entities"]}
        self.assertEqual("person", by_key["nguyen van an"]["kind"])
        self.assertEqual(2, by_key["nguyen van an"]["count"])
        self.assertTrue(profile["keywords"])

    def test_chunks_carry_entities_and_profile_summary(self):
        document = official_document()
        profile = build_document_profile(document)
        chunks = StructureAwareChunker(PipelineConfig().chunking).chunk(document)
        attach_profile_to_chunks(chunks, profile)
        body = next(chunk for chunk in chunks if "Dự án cầu" in chunk.content)
        self.assertIn("Nguyễn Văn An", body.entities)
        self.assertEqual("123/QĐ-UBND", body.payload()["doc_number"])
        self.assertIn("Trần Thị Mai", body.payload()["signers"])
        sparse = tokenize_vi(bm25_text(body))
        self.assertIn("nguyen", sparse)  # diacritic-free entity tokens for OCR-stripped queries

    def test_entity_lookup_ignores_diacritics_and_case(self):
        document = official_document()
        document.metadata["profile"] = build_document_profile(document)
        chunks = StructureAwareChunker(PipelineConfig().chunking).chunk(document)
        store = MetadataStore(Path(":memory:"))
        try:
            store.upsert_document(document, chunks)
            matches = store.match_entities("quyet dinh nao do NGUYEN VAN AN ky?")
            self.assertEqual({"doc1"}, {match["document_id"] for match in matches})
            self.assertEqual([], store.match_entities("khong lien quan"))
            self.assertEqual("123/QĐ-UBND", store.document_profile("doc1")["doc_number"])
            store.delete_document("doc1")
            self.assertEqual([], store.match_entities("nguyen van an"))
        finally:
            store.close()

    def test_neighbor_chunks_stay_inside_the_section(self):
        document = ParsedDocument("doc1", "a.pdf", "pdf", "hash")
        long_text = " ".join(f"Câu số {i} nói về quy trình phê duyệt hồ sơ." for i in range(60))
        document.blocks = [
            block("p1", "text", long_text, 0, section_path=["A"], page=1),
            block("p2", "text", "Mục khác.", 1, section_path=["B"], page=1),
        ]
        config = PipelineConfig()
        config.chunking.text_chunk_chars = 400
        chunks = StructureAwareChunker(config.chunking).chunk(document)
        store = MetadataStore(Path(":memory:"))
        try:
            store.upsert_document(document, chunks)
            middle = chunks[2]
            neighbors = store.neighbor_chunks(middle, 1)
            self.assertEqual([middle.ordinal - 1, middle.ordinal + 1], [c.ordinal for c in neighbors])
            last_a = next(c for c in chunks if c.section_path == ["A"] and c.ordinal == max(
                x.ordinal for x in chunks if x.section_path == ["A"]))
            self.assertTrue(all(c.section_path == ["A"] for c in store.neighbor_chunks(last_a, 2)))
        finally:
            store.close()

    def test_normalize_key(self):
        self.assertEqual("nguyen van an", normalize_key("  Nguyễn  Văn An "))
        self.assertEqual("dong nai", normalize_key("Đồng Nai"))


class FakeSentenceModel:
    def encode(self, texts, batch_size=8, normalize_embeddings=True, show_progress_bar=False):
        import numpy as np

        vectors = np.zeros((len(texts), 128), dtype=np.float32)
        for row, text in enumerate(texts):
            for token in tokenize_vi(text):
                vectors[row, hash(token) % 128] += 1.0
        norms = np.linalg.norm(vectors, axis=1, keepdims=True)
        return vectors / np.maximum(norms, 1e-9)


@unittest.skipUnless(has("docx", "openpyxl", "qdrant_client", "rank_bm25", "numpy"), "index deps missing")
class StructureAwareEndToEndTests(unittest.TestCase):
    def setUp(self):
        from rag_kaggle import RAGPipeline

        self.tmp = tempfile.TemporaryDirectory()
        root = Path(self.tmp.name)
        self.inputs = root / "inputs"
        self.inputs.mkdir()
        config = PipelineConfig(work_dir=root / "runtime")
        config.parsing.enable_ocr = False
        config.retrieval.reranker_enabled = False
        config.runtime.unload_models_between_stages = False
        config.chunking.table_inline_max_chars = 800
        config.chunking.text_chunk_chars = 300
        self.pipeline = RAGPipeline(config)
        self.pipeline.encoder.model = FakeSentenceModel()
        self.pipeline.encoder.model_name = "fake-bow"
        self.pipeline.generator._chat = lambda *a, **k: json.dumps(
            {"answer": "Theo tài liệu [SOURCE_1].", "source_ids": ["SOURCE_1"]}, ensure_ascii=False
        )

    def tearDown(self):
        self.pipeline.close()
        self.tmp.cleanup()

    def _docx(self):
        from docx import Document

        document = Document()
        document.add_heading("Báo cáo xã Hòa Bình", level=1)
        document.add_paragraph("Ông Nguyễn Văn An phụ trách dự án. " * 12)
        document.add_paragraph("Bảng 1: Doanh thu theo tháng", style="Caption")
        table = document.add_table(rows=31, cols=4)
        for r, row in enumerate(table.rows):
            for c, cell in enumerate(row.cells):
                cell.text = ["Tháng", "Doanh thu", "Chi phí", "Lãi"][c] if r == 0 else f"T{r}-{c}-{r * c * 1000}"
        path = self.inputs / "bao_cao.docx"
        document.save(path)
        return path

    def test_large_table_workbook_entities_and_full_table_context(self):
        import openpyxl

        report = self.pipeline.ingest([self._docx()])
        self.assertTrue(report["ok"], report["errors"])
        document_report = report["documents"][0]
        self.assertGreater(document_report["entities"], 0)

        chunks = self.pipeline.metadata.list_chunks()
        table_chunks = [chunk for chunk in chunks if chunk.chunk_type == "table"]
        self.assertEqual(1, len(table_chunks))  # a table is never cut into pieces
        table_chunk = table_chunks[0]
        self.assertTrue(table_chunk.metadata["table_truncated"])
        workbook_path = self.pipeline.config.work_dir / table_chunk.metadata["table_file"]
        self.assertTrue(workbook_path.exists())
        sheet = openpyxl.load_workbook(workbook_path).active
        self.assertEqual(31, sheet.max_row)
        self.assertEqual("T30-3-90000", sheet.cell(31, 4).value)
        self.assertTrue(all(chunk.ordinal == i for i, chunk in enumerate(chunks)))

        hits = self.pipeline.retriever.retrieve("Doanh thu theo tháng", top_k=5)
        hit = next(h for h in hits if h.chunk.chunk_type == "table")
        context = next(c for c in self.pipeline.retriever.expand_context([hit]))
        self.assertIn("T30-3-90000", context["content"])  # preview expanded to the full table
        self.assertTrue(context["table_file"].endswith(".xlsx"))

        matches = self.pipeline.metadata.match_entities("ông nguyen van an phụ trách gì?")
        self.assertTrue(matches)
        result = self.pipeline.ask("ông nguyen van an phụ trách gì?")
        trace = result["trace"]
        self.assertTrue(trace)
        self.assertTrue(result["contexts"])

    def test_entity_boost_adds_scoped_rankings(self):
        self.pipeline.ingest([self._docx()])
        _, stage_trace = self.pipeline.retriever.retrieve_with_trace(
            "Nguyen Van An", boost_document_ids=[d["document_id"] for d in self.pipeline.metadata.list_documents()]
        )
        scopes = {ranking["scope"] for ranking in stage_trace["rankings"]}
        self.assertEqual({"all", "entity"}, scopes)

    def test_legacy_schema_is_rejected_with_a_clear_message(self):
        import sqlite3

        legacy = Path(self.tmp.name) / "legacy.db"
        connection = sqlite3.connect(legacy)
        connection.execute("CREATE TABLE chunks (chunk_id TEXT, parent_id TEXT)")
        connection.commit()
        connection.close()
        with self.assertRaises(RuntimeError) as raised:
            MetadataStore(legacy)
        self.assertIn("legacy parent-child schema", str(raised.exception))


if __name__ == "__main__":
    unittest.main()
