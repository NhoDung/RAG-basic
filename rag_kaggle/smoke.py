"""End-to-end smoke test for a real Kaggle GPU session.

Runs every stage (environment, OCR, parse/chunk/index, embedding, reranker, LLM,
structured computation) on small generated documents in a separate work_dir and
reports PASS/WARN/FAIL per stage, so failures surface before using the demo.
"""

from __future__ import annotations

import copy
import importlib.metadata
import platform
import shutil
import time
import traceback
from pathlib import Path
from typing import Any, Callable

from .config import PipelineConfig


FONT_CANDIDATES = (
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
    "/usr/share/fonts/truetype/noto/NotoSans-Regular.ttf",
)
PACKAGES = (
    "torch",
    "transformers",
    "sentence-transformers",
    "FlagEmbedding",
    "bitsandbytes",
    "accelerate",
    "qdrant-client",
    "gradio",
    "pymupdf",
    "python-docx",
    "openpyxl",
)


def run_smoke_test(
    config: PipelineConfig,
    work_dir: str | Path | None = None,
    keep: bool = False,
    printer: Callable[[str], None] = print,
) -> dict[str, Any]:
    from .pipeline import RAGPipeline

    smoke_config = copy.deepcopy(config)
    smoke_config.work_dir = Path(work_dir or config.work_dir.parent / "rag_smoke")
    if smoke_config.work_dir.exists():
        shutil.rmtree(smoke_config.work_dir)
    smoke_config.runtime.unload_models_between_stages = True

    results: list[dict[str, Any]] = []
    state: dict[str, Any] = {}

    def step(name: str, func: Callable[[], tuple[str, Any]], requires: tuple[str, ...] = ()) -> None:
        missing = [item for item in requires if state.get(item) != "PASS"]
        if missing:
            results.append({"step": name, "status": "SKIP", "seconds": 0, "detail": f"requires {', '.join(missing)}"})
            printer(f"[SKIP] {name}: requires {', '.join(missing)}")
            return
        started = time.perf_counter()
        try:
            status, detail = func()
        except Exception as exc:
            status, detail = "FAIL", f"{type(exc).__name__}: {exc}\n{traceback.format_exc(limit=3)}"
        seconds = round(time.perf_counter() - started, 1)
        state[name] = status
        results.append({"step": name, "status": status, "seconds": seconds, "detail": detail})
        printer(f"[{status}] {name} ({seconds}s): {_short(detail)}")

    step("environment", lambda: _check_environment(smoke_config))
    pipeline = RAGPipeline(smoke_config)
    samples = make_sample_documents(smoke_config.work_dir / "samples")
    try:
        if smoke_config.parsing.enable_ocr:
            step("ocr", lambda: _check_ocr(pipeline, samples["ocr_image"]))
        step("ingest", lambda: _check_ingest(pipeline, samples, state.get("ocr") == "PASS"))
        step("dense_index", lambda: _check_index(pipeline), requires=("ingest",))
        step("retrieve_rerank", lambda: _check_retrieve(pipeline), requires=("dense_index",))
        step("llm_answer", lambda: _check_answer(pipeline), requires=("retrieve_rerank",))
        step("structured_computation", lambda: _check_computation(pipeline), requires=("llm_answer",))
        step("gpu_memory", _check_gpu_memory)
    finally:
        for model in (pipeline.generator, pipeline.reranker, pipeline.encoder, pipeline.ocr, pipeline.vision):
            try:
                model.unload()
            except Exception:
                pass
        pipeline.close()
        if not keep:
            shutil.rmtree(smoke_config.work_dir, ignore_errors=True)

    summary = {
        "ok": all(item["status"] in ("PASS", "WARN") for item in results),
        "steps": results,
    }
    printer("\n" + format_table(results))
    return summary


# ----------------------------------------------------------------- steps


def _check_environment(config: PipelineConfig) -> tuple[str, Any]:
    import torch

    detail: dict[str, Any] = {"python": platform.python_version(), "packages": {}}
    for package in PACKAGES:
        try:
            detail["packages"][package] = importlib.metadata.version(package)
        except importlib.metadata.PackageNotFoundError:
            detail["packages"][package] = None
    if not torch.cuda.is_available():
        return "FAIL", {**detail, "error": "No CUDA GPU visible to PyTorch. Select a T4 GPU accelerator."}
    devices = []
    for index in range(torch.cuda.device_count()):
        properties = torch.cuda.get_device_properties(index)
        devices.append(
            {
                "name": properties.name,
                "compute_capability": f"{properties.major}.{properties.minor}",
                "vram_gb": round(properties.total_memory / 1024**3, 1),
            }
        )
    detail.update({"torch_cuda": torch.version.cuda, "gpus": devices})
    missing = [name for name, version in detail["packages"].items() if version is None]
    major = torch.cuda.get_device_properties(0).major
    if major < 7:
        detail["warning"] = "GPU compute capability < 7.0 (e.g. P100): PaddleOCR-VL is not supported; use T4."
        return "WARN", detail
    if missing:
        detail["warning"] = f"Missing packages: {missing}"
        return "WARN", detail
    return "PASS", detail


def _check_ocr(pipeline, image_path: Path) -> tuple[str, Any]:
    result = pipeline.ocr.predict(image_path)
    text = result.get("text", "")
    detail = {
        "mode": "worker" if pipeline.config.parsing.ocr_python else "in-process",
        "chars": len(text),
        "layout_blocks": [block["label"] for block in result.get("blocks", [])][:10],
        "preview": text[:200],
    }
    pipeline.ocr.unload()
    if not text.strip():
        return "FAIL", {**detail, "error": "OCR returned no text"}
    expected = "499"
    return ("PASS" if expected in text else "WARN"), detail


def _check_ingest(pipeline, samples: dict[str, Path], ocr_ok: bool) -> tuple[str, Any]:
    files = [samples["docx"], samples["xlsx"], samples["pdf"], samples["scan_pdf"]]
    report = pipeline.ingest(files, reset=True)
    by_file = {item["file"]: item for item in report["documents"]}
    detail = {
        "documents": {name: {k: item[k] for k in ("blocks", "chunks", "relationships")} for name, item in by_file.items()},
        "errors": report["errors"],
        "stats": report["stats"],
    }
    core = [samples["docx"].name, samples["xlsx"].name, samples["pdf"].name]
    if any(by_file.get(name, {}).get("chunks", 0) == 0 for name in core):
        return "FAIL", detail
    if report["errors"]:
        return "FAIL", detail
    scan = by_file.get(samples["scan_pdf"].name, {})
    if pipeline.config.parsing.enable_ocr and ocr_ok and not scan.get("chunks"):
        detail["warning"] = "Scanned PDF produced no chunks"
        return "WARN", detail
    return "PASS", detail


def _check_index(pipeline) -> tuple[str, Any]:
    model = pipeline.metadata.get_setting("index.dense_model")
    dimension = pipeline.metadata.get_setting("index.dimension")
    status = "PASS" if model and dimension else "FAIL"
    return status, {"dense_model": model, "dimension": dimension, "chunks": pipeline.metadata.stats()["chunks"]}


def _check_retrieve(pipeline) -> tuple[str, Any]:
    hits = pipeline.retriever.retrieve("Phí thường niên thẻ Visa Gold là bao nhiêu?")
    if not hits:
        return "FAIL", {"error": "no hits"}
    top = hits[0]
    detail = {
        "top_file": top.chunk.source_file,
        "top_type": top.chunk.chunk_type,
        "dense_rank": top.dense_rank,
        "sparse_rank": top.sparse_rank,
        "rerank_score": top.rerank_score,
        "reranker_enabled": pipeline.config.retrieval.reranker_enabled,
    }
    if pipeline.config.retrieval.reranker_enabled and top.rerank_score is None:
        return "FAIL", {**detail, "error": "reranker produced no score"}
    found = any("499" in hit.chunk.content for hit in hits[:3])
    return ("PASS" if found else "WARN"), detail


def _check_answer(pipeline) -> tuple[str, Any]:
    result = pipeline.ask("Phí thường niên thẻ Visa Gold là bao nhiêu?")
    detail = {
        "answer": result["answer"][:300],
        "citations": result.get("citations"),
        "citation_valid": result.get("citation_valid"),
        "refused": result.get("refused"),
        "warnings": result.get("warnings"),
        "latency_ms": result.get("latency_ms"),
        "usage": pipeline.generator.last_usage,
    }
    if not result["answer"].strip() or result.get("refused"):
        return "FAIL", detail
    if "499" not in result["answer"] or not result.get("citation_valid"):
        return "WARN", detail
    return "PASS", detail


def _check_computation(pipeline) -> tuple[str, Any]:
    result = pipeline.ask("Tổng doanh thu khu vực Bắc là bao nhiêu?")
    computed = result.get("computation") or []
    detail = {
        "computed": [{"value": item["value"], "column": item["column"], "filters": item["filters"]} for item in computed],
        "answer": result["answer"][:300],
    }
    if not computed or abs(computed[0]["value"] - 200) > 1e-6:
        return "FAIL", detail
    return ("PASS" if "200" in result["answer"] else "WARN"), detail


def _check_gpu_memory() -> tuple[str, Any]:
    import torch

    if not torch.cuda.is_available():
        return "WARN", "no GPU"
    return "PASS", {
        "peak_allocated_gb": round(torch.cuda.max_memory_allocated() / 1024**3, 2),
        "currently_allocated_gb": round(torch.cuda.memory_allocated() / 1024**3, 2),
    }


# --------------------------------------------------------------- samples


def make_sample_documents(directory: Path) -> dict[str, Path]:
    """Small DOCX/XLSX/PDF files plus a scanned page; answers are known in advance."""
    directory.mkdir(parents=True, exist_ok=True)
    return {
        "docx": _sample_docx(directory / "smoke_bieu_phi.docx"),
        "xlsx": _sample_xlsx(directory / "smoke_doanh_thu.xlsx"),
        "pdf": _sample_pdf(directory / "smoke_dieu_kien.pdf"),
        **_sample_scan(directory),
    }


def _sample_docx(path: Path) -> Path:
    from docx import Document

    document = Document()
    document.add_heading("Biểu phí thẻ tín dụng", level=1)
    document.add_paragraph("Phí thường niên của từng loại thẻ được nêu tại Bảng 1.")
    document.add_paragraph("Bảng 1: Phí thường niên", style="Caption")
    table = document.add_table(rows=3, cols=2)
    for row, values in zip(table.rows, [["Loại thẻ", "Phí thường niên"], ["Visa Gold", "499.000 VNĐ"], ["Visa Platinum", "999.000 VNĐ"]]):
        for cell, value in zip(row.cells, values):
            cell.text = value
    document.save(path)
    return path


def _sample_xlsx(path: Path) -> Path:
    import openpyxl
    from openpyxl.chart import BarChart, Reference

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Doanh thu"
    sheet["A1"] = "Báo cáo doanh thu 2025 (tỷ đồng)"
    rows = [["Chi nhánh", "Khu vực", "Doanh thu"], ["Hà Nội", "Bắc", 120], ["Hải Phòng", "Bắc", 80], ["TP.HCM", "Nam", 300]]
    for offset, row in enumerate(rows, start=3):
        for column, value in enumerate(row, start=1):
            sheet.cell(offset, column, value)
    chart = BarChart()
    chart.title = "Doanh thu theo chi nhánh"
    chart.add_data(Reference(sheet, min_col=3, min_row=3, max_row=6), titles_from_data=True)
    chart.set_categories(Reference(sheet, min_col=1, min_row=4, max_row=6))
    sheet.add_chart(chart, "F2")
    workbook.save(path)
    return path


def _font_path() -> str | None:
    return next((path for path in FONT_CANDIDATES if Path(path).exists()), None)


def _sample_pdf(path: Path) -> Path:
    try:
        import pymupdf as fitz
    except ImportError:
        import fitz

    font = _font_path()
    pdf = fitz.open()
    page = pdf.new_page()
    kwargs = {"fontfile": font, "fontname": "smokefont"} if font else {}
    page.insert_text((72, 72), "1. Điều kiện mở thẻ", fontsize=18, **kwargs)
    y = 110
    for line in (
        "Khách hàng từ đủ 18 tuổi, có giấy tờ tùy thân hợp lệ.",
        "Thu nhập tối thiểu 10 triệu đồng mỗi tháng.",
        "Hồ sơ được thẩm định trong vòng 5 ngày làm việc.",
    ):
        page.insert_text((72, y), line, fontsize=11, **kwargs)
        y += 18
    pdf.save(path)
    pdf.close()
    return path


def _sample_scan(directory: Path) -> dict[str, Path]:
    """Render text into an image and embed it in a PDF page without a text layer."""
    try:
        import pymupdf as fitz
    except ImportError:
        import fitz
    from PIL import Image, ImageDraw, ImageFont

    font_path = _font_path()
    title_font = ImageFont.truetype(font_path, 44) if font_path else ImageFont.load_default()
    body_font = ImageFont.truetype(font_path, 32) if font_path else ImageFont.load_default()
    image = Image.new("RGB", (1240, 1754), "white")
    draw = ImageDraw.Draw(image)
    draw.text((100, 120), "BIỂU PHÍ DỊCH VỤ THẺ", fill="black", font=title_font)
    lines = [
        "Phí phát hành thẻ Visa Gold: 499.000 VNĐ",
        "Phí cấp lại mã PIN: 50.000 VNĐ",
        "Áp dụng từ ngày 01/01/2025.",
    ]
    for index, line in enumerate(lines):
        draw.text((100, 260 + index * 70), line, fill="black", font=body_font)
    image_path = directory / "smoke_scan_page.png"
    image.save(image_path)

    pdf_path = directory / "smoke_scan.pdf"
    pdf = fitz.open()
    page = pdf.new_page(width=595, height=842)
    page.insert_image(page.rect, filename=str(image_path))
    pdf.save(pdf_path)
    pdf.close()
    return {"scan_pdf": pdf_path, "ocr_image": image_path}


# --------------------------------------------------------------- output


def _short(detail: Any, limit: int = 220) -> str:
    text = str(detail).replace("\n", " ")
    return text if len(text) <= limit else text[:limit] + "..."


def format_table(results: list[dict[str, Any]]) -> str:
    lines = [f"{'STEP':<24}{'STATUS':<8}{'SECONDS':>8}", "-" * 40]
    lines.extend(f"{item['step']:<24}{item['status']:<8}{item['seconds']:>8}" for item in results)
    return "\n".join(lines)
