from __future__ import annotations

import datetime as dt
import json
import logging
import re
import time
import zipfile
from contextlib import contextmanager
from pathlib import Path
from typing import Callable, Iterable
from xml.etree import ElementTree

from .config import PARSER_VERSION, PipelineConfig
from .ingestion import IngestionError, convert_with_libreoffice, find_libreoffice
from .models import Block, ParsedDocument
from .paddleocr_vl import PaddleOCRVLAdapter
from .utils import (
    clean_text,
    excel_column_name,
    file_sha256,
    html_table_to_rows,
    markdown_table_to_rows,
    parse_cell_range,
    parse_number,
    rows_to_markdown,
    stable_id,
)
from .vision import VisionReasoner, classify_image, render_vlm_output


LOGGER = logging.getLogger(__name__)

# Re-exported for backwards compatibility with earlier imports.
__all__ = [
    "DocumentParser",
    "clean_text",
    "excel_column_name",
    "file_sha256",
    "rows_to_markdown",
    "save_parsed_document",
    "stable_id",
]

CAPTION_PATTERN = re.compile(
    r"^(hình|bảng|biểu đồ|sơ đồ|lưu đồ|đồ thị|figure|fig\.|table|chart)\s*(\d+(?:[.\-]\d+)*)\s*[:.\-–]?",
    re.IGNORECASE,
)
OCR_HEADING_LABELS = {"doc_title", "paragraph_title", "title", "header_title", "section_title"}
OCR_TABLE_LABELS = {"table"}
OCR_IMAGE_LABELS = {"image", "figure", "chart", "seal"}
OCR_SKIP_LABELS = {"header", "footer", "number", "page_number", "footnote_number"}

DRAWING_NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
}
R_EMBED = f"{{{DRAWING_NS['r']}}}embed"
R_ID = f"{{{DRAWING_NS['r']}}}id"


class DocumentParser:
    def __init__(
        self,
        config: PipelineConfig,
        ocr: PaddleOCRVLAdapter | None = None,
        vision: VisionReasoner | None = None,
    ):
        self.config = config
        self.ocr = ocr
        self.vision = vision
        self._progress: Callable[[str], None] | None = None

    def parse(
        self,
        path: str | Path,
        display_name: str | None = None,
        progress: Callable[[str], None] | None = None,
    ) -> ParsedDocument:
        source = Path(path)
        extension = source.suffix.lower()
        previous_progress, self._progress = self._progress, progress
        try:
            if extension == ".pdf":
                document = self._parse_pdf(source)
            elif extension == ".docx":
                document = self._parse_docx(source)
            elif extension in {".xlsx", ".xlsm", ".xltx", ".xltm"}:
                document = self._parse_excel(source)
            elif extension == ".xls":
                document = self._parse_legacy_excel(source)
            else:
                raise IngestionError("unsupported_type", f"Unsupported file type: {extension}")
        finally:
            self._progress = previous_progress
        if display_name:
            document.source_file = display_name
            for block in document.blocks:
                block.source_file = display_name
        return document

    def _notify(self, message: str) -> None:
        if self._progress is not None:
            self._progress(message)

    @contextmanager
    def _timed_unit(self, document: ParsedDocument, operation: str, unit: str, **details):
        started_at = dt.datetime.now(dt.timezone.utc)
        started_perf = time.perf_counter()
        status = "success"
        self._notify(f"[TIMING START] {operation} | {unit} | {started_at.isoformat(timespec='seconds')}")
        try:
            yield
        except Exception:
            status = "failed"
            raise
        finally:
            finished_at = dt.datetime.now(dt.timezone.utc)
            elapsed_seconds = round(time.perf_counter() - started_perf, 3)
            timing = {
                "operation": operation,
                "unit": unit,
                "started_at": started_at.isoformat(timespec="seconds"),
                "finished_at": finished_at.isoformat(timespec="seconds"),
                "elapsed_seconds": elapsed_seconds,
                "status": status,
                **details,
            }
            document.metadata.setdefault("timings", []).append(timing)
            self._notify(
                f"[TIMING END] {operation} | {unit} | {finished_at.isoformat(timespec='seconds')} "
                f"| elapsed_seconds={elapsed_seconds:.3f} | status={status}"
            )

    # ------------------------------------------------------------------ common

    def _new_document(self, source: Path) -> ParsedDocument:
        digest = file_sha256(source)
        document_id = stable_id(source.name, digest)
        return ParsedDocument(
            document_id=document_id,
            source_file=source.name,
            file_type=source.suffix.lower().lstrip("."),
            content_hash=digest,
            parser_version=PARSER_VERSION,
            created_at=dt.datetime.now(dt.timezone.utc).isoformat(timespec="seconds"),
            metadata={"warnings": []},
        )

    def _asset_dir(self, document_id: str) -> Path:
        target = self.config.asset_dir / document_id
        target.mkdir(parents=True, exist_ok=True)
        return target

    def _artifact_ref(self, path: Path) -> str:
        """Store corpus-relative paths so bundles can move between Kaggle sessions."""
        try:
            return path.resolve().relative_to(self.config.work_dir.resolve()).as_posix()
        except ValueError:
            return str(path)

    def _block(self, document: ParsedDocument, block_type: str, content: str, ordinal: object, **kwargs) -> Block:
        return Block(
            block_id=stable_id(document.document_id, block_type, ordinal),
            document_id=document.document_id,
            block_type=block_type,
            content=clean_text(content),
            source_file=document.source_file,
            **kwargs,
        )

    def _warn(self, document: ParsedDocument, stage: str, message: str, **extra) -> None:
        document.metadata.setdefault("warnings", []).append({"stage": stage, "message": message, **extra})

    def _add_table(self, document, rows, ordinal, **kwargs) -> Block | None:
        rows = [[clean_text(str(cell)) for cell in row] for row in rows if any(str(cell).strip() for cell in row)]
        if not rows:
            return None
        metadata = kwargs.pop("metadata", {})
        metadata["rows"] = rows
        block = self._block(document, "table", rows_to_markdown(rows), ordinal, metadata=metadata, **kwargs)
        return document.add_block(block)

    def _add_visual(
        self,
        document: ParsedDocument,
        image_path: Path,
        ordinal: str,
        section_path: list[str],
        caption: str = "",
        **kwargs,
    ) -> Block | None:
        """OCR an image, classify it and optionally enrich it with a VLM block."""
        if not self._image_is_large_enough(image_path):
            return None  # Logos/icons are treated as decorative and not indexed.
        location = (
            f"PDF page {kwargs['page']} image {image_path.name}"
            if kwargs.get("page") is not None
            else f"Excel sheet {kwargs['sheet_name']} image {image_path.name}"
            if kwargs.get("sheet_name")
            else f"image {image_path.name}"
        )
        ocr_result = self._run_ocr(document, image_path, unit=location)
        labels = [block["label"] for block in ocr_result.get("blocks", [])]
        image_type = classify_image(caption, ocr_result["text"], labels)
        metadata = kwargs.pop("metadata", {})
        metadata.update({"ocr": ocr_result["raw"], "image_type": image_type})

        vlm_result = {"status": "skipped", "output": None}
        if self.vision is not None and image_type in ("flowchart", "chart"):
            with self._timed_unit(document, "vlm", location, asset=str(image_path), image_type=image_type):
                vlm_result = self.vision.describe(
                    image_path,
                    image_type,
                    ocr_text=ocr_result["text"],
                    caption=caption,
                    section=" > ".join(section_path),
                )
            if vlm_result["status"] not in ("success", "skipped"):
                self._warn(document, "vlm", f"VLM {vlm_result['status']} for {image_path.name}")

        if not ocr_result["text"] and vlm_result["output"] is None:
            return None
        image_block = document.add_block(
            self._block(
                document,
                "image",
                ocr_result["text"] or "(Ảnh không có chữ nhận dạng được)",
                ordinal,
                section_path=list(section_path),
                asset_path=self._artifact_ref(image_path),
                metadata=metadata,
                raw_content=ocr_result.get("blocks") or None,
                **kwargs,
            )
        )
        if vlm_result["output"] is not None:
            document.add_block(
                self._block(
                    document,
                    image_type,
                    render_vlm_output(vlm_result["output"]),
                    f"{ordinal}-vlm",
                    section_path=list(section_path),
                    asset_path=self._artifact_ref(image_path),
                    metadata={"derived_from": image_block.block_id, "vlm_model": self.config.vision.model},
                    raw_content=vlm_result["output"],
                    page=kwargs.get("page"),
                    sheet_name=kwargs.get("sheet_name"),
                )
            )
        elif vlm_result["status"] == "needs_review":
            image_block.metadata["vlm_status"] = "needs_review"
        return image_block

    def _run_ocr(self, document: ParsedDocument, image_path: Path, unit: str | None = None) -> dict:
        if not self.config.parsing.enable_ocr or self.ocr is None:
            return {"text": "", "blocks": [], "raw": [], "model": None}
        try:
            with self._timed_unit(document, "ocr", unit or image_path.name, asset=str(image_path)):
                return self.ocr.predict(image_path)
        except Exception as exc:
            LOGGER.exception("OCR failed for %s", image_path)
            self._warn(document, "ocr", f"OCR failed for {image_path.name}: {exc}", asset=str(image_path))
            return {"text": "", "blocks": [], "raw": [], "model": self.config.parsing.ocr_model_name, "error": str(exc)}

    def _image_is_large_enough(self, path: Path) -> bool:
        try:
            from PIL import Image

            with Image.open(path) as image:
                return (
                    image.width >= self.config.parsing.min_image_width
                    and image.height >= self.config.parsing.min_image_height
                )
        except Exception:
            return False

    # --------------------------------------------------------------------- PDF

    def _parse_pdf(self, source: Path) -> ParsedDocument:
        try:
            import pymupdf as fitz
        except ImportError:
            import fitz

        document = self._new_document(source)
        asset_dir = self._asset_dir(document.document_id)
        pdf = fitz.open(source)
        if pdf.needs_pass:
            pdf.close()
            raise IngestionError("encrypted_file", f"{source.name} is password protected")

        body_size, heading_levels = _pdf_font_profile(pdf, self.config.parsing.heading_font_ratio)
        document.metadata.update({"pages": pdf.page_count, "body_font_size": body_size})
        sections: list[str] = []

        for page_index, page in enumerate(pdf):
            page_number = page_index + 1
            unit = f"PDF page {page_number}/{pdf.page_count}"
            with self._timed_unit(document, "parse", unit, page=page_number):
                page_text = clean_text(page.get_text("text"))
                if len(page_text) < self.config.parsing.scan_text_threshold:
                    if self.config.parsing.enable_ocr:
                        self._notify(f"{unit}: rendering and PaddleOCR-VL.")
                        sections = self._parse_scanned_page(document, page, page_number, asset_dir, sections)
                    elif page_text:
                        self._notify(f"{unit}: short text, OCR disabled.")
                        document.add_block(
                            self._block(document, "text", page_text, f"page-{page_number}-raw", page=page_number)
                        )
                else:
                    self._notify(f"{unit}: extracting text, tables, and images.")
                    sections = self._parse_digital_page(
                        document, pdf, page, page_number, asset_dir, sections, body_size, heading_levels
                    )

        pdf.close()
        return document

    def _parse_digital_page(self, document, pdf, page, page_number, asset_dir, sections, body_size, heading_levels):
        items = []  # (y0, x0, kind, payload)
        table_boxes = []
        try:
            tables = list(getattr(page.find_tables(), "tables", []))
        except (AttributeError, RuntimeError, ValueError):
            tables = []
        for table_index, table in enumerate(tables):
            rows = table.extract() or []
            normalized = [["" if cell is None else clean_text(str(cell)) for cell in row] for row in rows]
            if not any(any(row) for row in normalized):
                continue
            bbox = [float(value) for value in table.bbox]
            table_boxes.append(bbox)
            items.append((bbox[1], bbox[0], "table", (table_index, normalized, bbox)))

        for block_index, raw in enumerate(page.get_text("dict").get("blocks", [])):
            if raw.get("type") != 0:
                continue
            bbox = [float(value) for value in raw["bbox"]]
            if any(_bbox_inside(bbox, box) for box in table_boxes):
                continue  # Text already captured as table rows.
            lines, sizes, bold_chars, total_chars = [], [], 0, 0
            for line in raw.get("lines", []):
                spans = line.get("spans", [])
                lines.append("".join(span.get("text", "") for span in spans))
                for span in spans:
                    count = len(span.get("text", "").strip())
                    total_chars += count
                    sizes.extend([round(span.get("size", 0), 1)] * max(count, 1))
                    if span.get("flags", 0) & 16:
                        bold_chars += count
            text = clean_text("\n".join(lines))
            if not text:
                continue
            items.append(
                (
                    bbox[1],
                    bbox[0],
                    "text",
                    (block_index, text, bbox, max(sizes) if sizes else body_size, bold_chars / max(total_chars, 1)),
                )
            )

        if self.config.parsing.ocr_images_in_digital_docs:
            for image_index, image_info in enumerate(page.get_images(full=True)):
                rects = page.get_image_rects(image_info[0]) if hasattr(page, "get_image_rects") else []
                bbox = [float(v) for v in rects[0]] if rects else [0.0, 0.0, 0.0, 0.0]
                items.append((bbox[1], bbox[0], "image", (image_index, image_info[0], bbox)))

        # Keep PyMuPDF's native reading order for text (multi-column safe) and insert
        # tables/images before the first text block that starts below them.
        text_items = [item for item in items if item[2] == "text"]
        for item in sorted((item for item in items if item[2] != "text"), key=lambda value: value[0]):
            position = next(
                (index for index, existing in enumerate(text_items) if existing[2] == "text" and existing[0] >= item[0]),
                len(text_items),
            )
            text_items.insert(position, item)
        items = text_items
        last_caption = ""
        for _, _, kind, payload in items:
            if kind == "table":
                table_index, rows, bbox = payload
                self._add_table(
                    document,
                    rows,
                    f"page-{page_number}-table-{table_index}",
                    section_path=list(sections),
                    page=page_number,
                    bbox=bbox,
                )
            elif kind == "text":
                block_index, text, bbox, size, bold_ratio = payload
                level = _heading_level(
                    text, size, bold_ratio, body_size, heading_levels, self.config.parsing.heading_max_chars
                )
                if CAPTION_PATTERN.match(text) and len(text) <= 300:
                    block_type = "caption"
                    last_caption = text
                elif level:
                    block_type = "heading"
                    sections = sections[: level - 1] + [text.replace("\n", " ")]
                else:
                    block_type = "text"
                document.add_block(
                    self._block(
                        document,
                        block_type,
                        text,
                        f"page-{page_number}-text-{block_index}",
                        section_path=list(sections),
                        page=page_number,
                        bbox=bbox,
                        metadata={"font_size": size, "heading_level": level} if level else {"font_size": size},
                    )
                )
            else:
                image_index, xref, bbox = payload
                try:
                    extracted = pdf.extract_image(xref)
                except (RuntimeError, ValueError):
                    continue
                if not extracted or not extracted.get("image"):
                    continue
                width, height = int(extracted.get("width", 0)), int(extracted.get("height", 0))
                if width < self.config.parsing.min_image_width or height < self.config.parsing.min_image_height:
                    continue
                image_path = asset_dir / f"page_{page_number:04d}_image_{image_index:03d}.{extracted.get('ext', 'png')}"
                image_path.write_bytes(extracted["image"])
                self._add_visual(
                    document,
                    image_path,
                    f"page-{page_number}-image-{image_index}",
                    sections,
                    caption=last_caption,
                    page=page_number,
                    bbox=bbox,
                    metadata={"width": width, "height": height},
                )
        return sections

    def _parse_scanned_page(self, document, page, page_number, asset_dir, sections):
        image_path = asset_dir / f"page_{page_number:04d}.png"
        with self._timed_unit(document, "render", f"PDF page {page_number}", page=page_number):
            pixmap = page.get_pixmap(dpi=self.config.parsing.render_dpi, alpha=False)
            pixmap.save(image_path)
        ocr_result = self._run_ocr(document, image_path, unit=f"PDF page {page_number}")
        layout_blocks = ocr_result.get("blocks") or []
        common = {"page": page_number, "asset_path": self._artifact_ref(image_path)}

        if not layout_blocks:
            if ocr_result["text"]:
                document.add_block(
                    self._block(
                        document,
                        "ocr_page",
                        ocr_result["text"],
                        f"page-{page_number}",
                        section_path=list(sections),
                        metadata={"ocr": ocr_result["raw"], "is_scan": True},
                        raw_content=ocr_result["raw"],
                        **common,
                    )
                )
            elif "error" in ocr_result:
                self._warn(document, "ocr", f"Scanned page {page_number} has no OCR output", page=page_number)
            return sections

        # Raw OCR output is kept once per page for audit/re-parsing.
        document.metadata.setdefault("ocr_raw_pages", {})[str(page_number)] = ocr_result["raw"]
        for index, item in enumerate(layout_blocks):
            label, content = item["label"], item["content"]
            if label in OCR_SKIP_LABELS:
                continue
            ordinal = f"page-{page_number}-ocr-{index}"
            extra = {"bbox": item["bbox"], "confidence": item["confidence"], "metadata": {"ocr_label": label, "is_scan": True}}
            if label in OCR_TABLE_LABELS:
                rows = html_table_to_rows(content) if "<t" in content.lower() else markdown_table_to_rows(content)
                if rows:
                    self._add_table(document, rows, ordinal, section_path=list(sections), **common, **extra)
                    continue
            if label in OCR_HEADING_LABELS:
                text = clean_text(re.sub(r"^#+\s*", "", content))
                level = 1 if label == "doc_title" or not sections else 2
                sections = sections[: level - 1] + [text]
                block_type = "heading"
            elif CAPTION_PATTERN.match(content):
                block_type = "caption"
            elif label in OCR_IMAGE_LABELS:
                block_type = "chart" if label == "chart" else "image"
            else:
                block_type = "text"
            document.add_block(
                self._block(document, block_type, content, ordinal, section_path=list(sections), **common, **extra)
            )
        return sections

    # -------------------------------------------------------------------- DOCX

    def _parse_docx(self, source: Path) -> ParsedDocument:
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        document = self._new_document(source)
        docx = Document(source)
        asset_dir = self._asset_dir(document.document_id)
        sections: list[str] = []
        seen_parts: set[str] = set()
        last_caption = ""

        headers, footers = [], []
        for section in docx.sections:
            for container, target in ((section.header, headers), (section.footer, footers)):
                text = clean_text("\n".join(paragraph.text for paragraph in container.paragraphs))
                if text and text not in target:
                    target.append(text)
        document.metadata.update({"headers": headers, "footers": footers})

        body_items = list(iter_docx_blocks(docx))
        total_items = len(body_items)
        self._notify(f"DOCX: reading {total_items} body block(s), headers, tables, charts, and images.")
        for ordinal, item in enumerate(body_items, start=1):
            unit = f"DOCX block {ordinal}/{total_items}"
            with self._timed_unit(document, "parse", unit, block=ordinal):
                if ordinal == 1 or ordinal == total_items or ordinal % 10 == 0:
                    self._notify(f"{unit}.")
                if isinstance(item, Paragraph):
                    text = clean_text(item.text)
                    style_name = item.style.name if item.style is not None else ""
                    level = _docx_heading_level(item, style_name)
                    if text:
                        if level:
                            block_type = "heading"
                            sections = sections[: level - 1] + [text]
                        elif style_name.lower().startswith(("caption", "chú thích")) or (
                            CAPTION_PATTERN.match(text) and len(text) <= 300
                        ):
                            block_type = "caption"
                            last_caption = text
                        else:
                            block_type = "text"
                        document.add_block(
                            self._block(
                                document,
                                block_type,
                                text,
                                f"docx-{ordinal}",
                                section_path=list(sections),
                                metadata={"style": style_name, "heading_level": level},
                            )
                        )
                    self._extract_paragraph_drawings(
                        document, docx, item, ordinal, asset_dir, sections, seen_parts, last_caption
                    )
                elif isinstance(item, Table):
                    self._add_docx_table(document, item, f"docx-{ordinal}", sections)

        self._notify("DOCX: extracting media not anchored in the body flow.")
        self._extract_unreferenced_media(document, source, asset_dir, sections, seen_parts)
        self._notify("DOCX: extracting embedded workbooks.")
        self._extract_docx_embedded_workbooks(document, source)
        return document

    def _add_docx_table(self, document, table, ordinal, sections, depth=0):
        # python-docx repeats merged cells in every spanned column; keeping the repeated
        # value preserves column alignment (merged-cell normalization).
        rows = [[clean_text(cell.text) for cell in row.cells] for row in table.rows]
        block = self._add_table(
            document,
            rows,
            ordinal,
            section_path=list(sections),
            metadata={"nested_depth": depth} if depth else {},
        )
        for row_index, row in enumerate(table.rows):
            for cell_index, cell in enumerate(row.cells):
                for nested_index, nested in enumerate(cell.tables):
                    nested_block = self._add_docx_table(
                        document, nested, f"{ordinal}-nested-{row_index}-{cell_index}-{nested_index}", sections, depth + 1
                    )
                    if nested_block is not None and block is not None:
                        nested_block.metadata["nested_in"] = block.block_id
        return block

    def _extract_paragraph_drawings(self, document, docx, paragraph, ordinal, asset_dir, sections, seen, caption):
        element = paragraph._p
        part = docx.part
        for index, blip in enumerate(element.iter(f"{{{DRAWING_NS['a']}}}blip")):
            rel_id = blip.get(R_EMBED)
            image_part = part.related_parts.get(rel_id) if rel_id else None
            if image_part is None:
                continue
            partname = str(image_part.partname)
            seen.add(partname)
            suffix = Path(partname).suffix.lower() or ".png"
            if suffix not in {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff", ".gif"}:
                continue
            image_path = asset_dir / f"docx_{ordinal:04d}_image_{index:02d}{suffix}"
            image_path.write_bytes(image_part.blob)
            self._add_visual(
                document,
                image_path,
                f"docx-{ordinal}-image-{index}",
                sections,
                caption=caption,
                metadata={"archive_path": partname.lstrip("/"), "relationship_id": rel_id},
            )
        for index, chart_ref in enumerate(element.iter(f"{{{DRAWING_NS['c']}}}chart")):
            rel_id = chart_ref.get(R_ID)
            chart_part = part.related_parts.get(rel_id) if rel_id else None
            if chart_part is None:
                continue
            seen.add(str(chart_part.partname))
            chart = parse_chart_xml(chart_part.blob)
            if chart is None:
                continue
            document.add_block(
                self._block(
                    document,
                    "chart",
                    render_chart(chart, caption),
                    f"docx-{ordinal}-chart-{index}",
                    section_path=list(sections),
                    metadata={**chart, "relationship_id": rel_id, "caption": caption or None},
                )
            )

    def _extract_unreferenced_media(self, document, source, asset_dir, sections, seen):
        """Images outside the main body flow (text boxes, table cells) are processed last."""
        supported = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff"}
        with zipfile.ZipFile(source) as archive:
            media = [name for name in archive.namelist() if name.startswith("word/media/")]
            for index, name in enumerate(media):
                if f"/{name}" in seen or Path(name).suffix.lower() not in supported:
                    continue
                image_path = asset_dir / f"docx_media_{index:03d}{Path(name).suffix.lower()}"
                image_path.write_bytes(archive.read(name))
                self._add_visual(
                    document,
                    image_path,
                    f"docx-image-{index}",
                    sections,
                    metadata={"archive_path": name, "placement": "unreferenced"},
                )

    def _extract_docx_embedded_workbooks(self, document, source):
        with zipfile.ZipFile(source) as archive:
            workbook_names = [
                name
                for name in archive.namelist()
                if name.startswith("word/embeddings/") and Path(name).suffix.lower() in {".xlsx", ".xlsm"}
            ]
            for index, name in enumerate(workbook_names):
                embedded = self.config.source_dir / f"{document.document_id}_embedded_{index}.xlsx"
                embedded.write_bytes(archive.read(name))
                try:
                    parsed = self._parse_excel(embedded)
                except Exception as exc:  # Embedded chart workbooks are frequently partial.
                    self._warn(document, "parse", f"Could not parse embedded workbook {name}: {exc}")
                    continue
                for block in parsed.blocks:
                    block.document_id = document.document_id
                    block.source_file = document.source_file
                    block.block_id = stable_id(document.document_id, "embedded", index, block.block_id)
                    block.metadata["embedded_workbook"] = name
                    if block.metadata.get("derived_from"):
                        block.metadata["derived_from"] = stable_id(
                            document.document_id, "embedded", index, block.metadata["derived_from"]
                        )
                    document.add_block(block)

    # ------------------------------------------------------------------- Excel

    def _parse_excel(self, source: Path, display_source: Path | None = None) -> ParsedDocument:
        import openpyxl
        from openpyxl.utils import get_column_letter

        document = self._new_document(display_source or source)
        asset_dir = self._asset_dir(document.document_id)
        # Read twice: formulas and cached values (openpyxl does not recalculate).
        formulas = openpyxl.load_workbook(source, data_only=False)
        values = openpyxl.load_workbook(source, data_only=True)
        hidden_sheets, missing_cached = [], 0
        sheet_regions: dict[str, list[tuple[Block, tuple[int, int, int, int]]]] = {}

        total_sheets = len(formulas.worksheets)
        self._notify(f"Excel: reading {total_sheets} worksheet(s), cells, merged ranges, and formulas.")
        for sheet_index, formula_sheet in enumerate(formulas.worksheets):
            unit = f"Excel sheet {sheet_index + 1}/{total_sheets}: {formula_sheet.title}"
            with self._timed_unit(document, "parse", unit, sheet_name=formula_sheet.title):
                self._notify(f"{unit}.")
                if formula_sheet.sheet_state != "visible":
                    hidden_sheets.append(formula_sheet.title)
                    if not self.config.parsing.include_hidden_sheets:
                        continue
                value_sheet = values[formula_sheet.title]
                grid, formula_map, hidden = self._read_sheet_grid(formula_sheet, value_sheet)
                missing_cached += sum(1 for info in formula_map.values() if info["cached_value"] is None)
                title: str | None = None

                for region_index, (min_row, max_row, min_col, max_col) in enumerate(split_excel_regions_from_grid(grid)):
                    rows = [
                        [grid[r - 1][c - 1] if c - 1 < len(grid[r - 1]) else "" for c in range(min_col, max_col + 1)]
                        for r in range(min_row, max_row + 1)
                    ]
                    cell_range = f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max_row}"
                    region_formulas = {
                        coordinate: info
                        for coordinate, info in formula_map.items()
                        if min_row <= info["row"] <= max_row and min_col <= info["col"] <= max_col
                    }
                    kind = classify_excel_region(rows)
                    section = [formula_sheet.title] + ([title] if title and kind != "text" else [])
                    common = {
                        "section_path": section,
                        "sheet_name": formula_sheet.title,
                        "cell_range": cell_range,
                    }
                    region_meta = {
                        "formulas": region_formulas,
                        "origin": [min_row, min_col],
                        "numeric_repr": "python",
                    }
                    ordinal = f"sheet-{sheet_index}-region-{region_index}"
                    if kind == "text":
                        text = "\n".join(dict.fromkeys(cell for row in rows for cell in row if cell))
                        title = text.split("\n")[0][:160]
                        block = document.add_block(
                            self._block(document, "text", text, ordinal, metadata=region_meta, **common)
                        )
                    elif kind == "kpi":
                        block = document.add_block(
                            self._block(
                                document,
                                "kpi",
                                render_kpi(rows),
                                ordinal,
                                metadata={**region_meta, "rows": rows},
                                **common,
                            )
                        )
                    else:
                        block = self._add_table(document, rows, ordinal, metadata=region_meta, **common)
                    if block is not None:
                        sheet_regions.setdefault(formula_sheet.title, []).append(
                            (block, (min_row, max_row, min_col, max_col))
                        )
                if hidden["rows"] or hidden["columns"]:
                    document.metadata.setdefault("hidden_ranges", {})[formula_sheet.title] = hidden

        self._notify("Excel: extracting charts and embedded images.")
        for sheet_index, formula_sheet in enumerate(formulas.worksheets):
            if formula_sheet.title in hidden_sheets and not self.config.parsing.include_hidden_sheets:
                continue
            unit = f"Excel sheet {sheet_index + 1}/{total_sheets}: {formula_sheet.title}"
            with self._timed_unit(document, "extract_visuals", unit, sheet_name=formula_sheet.title):
                for chart_index, chart in enumerate(getattr(formula_sheet, "_charts", [])):
                    info = openpyxl_chart_info(chart, values)
                    document.add_block(
                        self._block(
                            document,
                            "chart",
                            render_chart(info),
                            f"sheet-{sheet_index}-chart-{chart_index}",
                            section_path=[formula_sheet.title],
                            sheet_name=formula_sheet.title,
                            metadata=info,
                        )
                    )
                for image_index, image in enumerate(getattr(formula_sheet, "_images", [])):
                    try:
                        image_bytes = image._data()
                    except (AttributeError, ValueError):
                        continue
                    extension = Path(getattr(image, "path", "image.png")).suffix or ".png"
                    image_path = asset_dir / f"sheet_{sheet_index}_image_{image_index}{extension}"
                    image_path.write_bytes(image_bytes)
                    self._add_visual(
                        document,
                        image_path,
                        f"sheet-{sheet_index}-image-{image_index}",
                        [formula_sheet.title],
                        sheet_name=formula_sheet.title,
                    )

        document.metadata.update(
            {
                "sheets": [sheet.title for sheet in formulas.worksheets],
                "hidden_sheets": hidden_sheets,
                "formula_cells_without_cached_value": missing_cached,
            }
        )
        if missing_cached:
            self._warn(
                document,
                "parse",
                f"{missing_cached} formula cells have no cached value; recalculate the workbook "
                "(LibreOffice/Excel) and re-ingest for exact numbers",
            )
        formulas.close()
        values.close()
        return document

    def _read_sheet_grid(self, formula_sheet, value_sheet):
        from openpyxl.utils import get_column_letter

        include_hidden = self.config.parsing.include_hidden_rows_columns
        hidden_rows = {
            index for index, dimension in formula_sheet.row_dimensions.items() if dimension.hidden
        }
        hidden_columns = set()
        for key, dimension in formula_sheet.column_dimensions.items():
            if dimension.hidden:
                start = dimension.min or 0
                end = dimension.max or start
                if start:
                    hidden_columns.update(range(start, end + 1))
                else:
                    from openpyxl.utils import column_index_from_string

                    hidden_columns.add(column_index_from_string(key))

        value_rows = list(value_sheet.iter_rows())
        formula_rows = list(formula_sheet.iter_rows())
        grid: list[list[str]] = []
        formula_map: dict[str, dict] = {}
        for row_offset, value_row in enumerate(value_rows):
            row_number = row_offset + 1
            formula_row = formula_rows[row_offset] if row_offset < len(formula_rows) else ()
            row_values = []
            for col_offset, value_cell in enumerate(value_row):
                col_number = col_offset + 1
                formula_cell = formula_row[col_offset] if col_offset < len(formula_row) else None
                formula = formula_cell.value if formula_cell is not None else None
                display = format_cell_value(value_cell.value, value_cell.number_format)
                if isinstance(formula, str) and formula.startswith("="):
                    formula_map[f"{get_column_letter(col_number)}{row_number}"] = {
                        "formula": formula,
                        "cached_value": value_cell.value if not isinstance(value_cell.value, (dt.date, dt.datetime)) else str(value_cell.value),
                        "number_format": value_cell.number_format,
                        "display_value": display,
                        "row": row_number,
                        "col": col_number,
                    }
                    if value_cell.value is None:
                        display = formula
                if not include_hidden and (row_number in hidden_rows or col_number in hidden_columns):
                    display = ""
                row_values.append(display)
            grid.append(row_values)

        for merged in value_sheet.merged_cells.ranges:
            top_value = grid[merged.min_row - 1][merged.min_col - 1] if merged.min_row <= len(grid) else ""
            for r in range(merged.min_row, merged.max_row + 1):
                for c in range(merged.min_col, merged.max_col + 1):
                    if r - 1 < len(grid) and c - 1 < len(grid[r - 1]) and not grid[r - 1][c - 1]:
                        grid[r - 1][c - 1] = top_value
        hidden = {
            "rows": sorted(hidden_rows),
            "columns": [get_column_letter(c) for c in sorted(hidden_columns)],
            "indexed": include_hidden,
        }
        return grid, formula_map, hidden

    def _parse_legacy_excel(self, source: Path) -> ParsedDocument:
        binary = find_libreoffice(self.config.parsing.libreoffice_binary)
        if binary:
            converted = convert_with_libreoffice(source, "xlsx", self.config.work_dir / "converted", binary)
            document = self._parse_excel(converted, display_source=source)
            document.file_type = "xls"
            document.metadata["converted_with"] = "libreoffice"
            return document

        import pandas as pd

        document = self._new_document(source)
        document.metadata["converted_with"] = "calamine"
        sheets = pd.read_excel(source, sheet_name=None, header=None, engine="calamine")
        for sheet_index, (sheet_name, frame) in enumerate(sheets.items()):
            unit = f"Legacy Excel sheet {sheet_index + 1}/{len(sheets)}: {sheet_name}"
            with self._timed_unit(document, "parse", unit, sheet_name=str(sheet_name)):
                frame = frame.dropna(how="all").dropna(axis=1, how="all")
                if frame.empty:
                    continue
                rows = [
                    ["" if pd.isna(value) else str(value) for value in row]
                    for row in frame.itertuples(index=False, name=None)
                ]
                first_row = int(frame.index[0]) + 1
                first_col = int(frame.columns[0]) + 1
                cell_range = (
                    f"{excel_column_name(first_col)}{first_row}:"
                    f"{excel_column_name(int(frame.columns[-1]) + 1)}{int(frame.index[-1]) + 1}"
                )
                self._add_table(
                    document,
                    rows,
                    f"legacy-sheet-{sheet_index}",
                    section_path=[str(sheet_name)],
                    sheet_name=str(sheet_name),
                    cell_range=cell_range,
                    metadata={"legacy_xls": True, "origin": [first_row, first_col], "numeric_repr": "python"},
                )
        return document


# ---------------------------------------------------------------------- helpers


def _bbox_inside(inner: list[float], outer: list[float], tolerance: float = 2.0) -> bool:
    x0, y0, x1, y1 = inner
    ox0, oy0, ox1, oy1 = outer
    cx, cy = (x0 + x1) / 2, (y0 + y1) / 2
    return ox0 - tolerance <= cx <= ox1 + tolerance and oy0 - tolerance <= cy <= oy1 + tolerance


def _pdf_font_profile(pdf, ratio: float) -> tuple[float, list[float]]:
    """Return the body font size and heading font sizes (largest first)."""
    counts: dict[float, int] = {}
    for page in pdf:
        for block in page.get_text("dict").get("blocks", []):
            for line in block.get("lines", []):
                for span in line.get("spans", []):
                    size = round(span.get("size", 0), 1)
                    counts[size] = counts.get(size, 0) + len(span.get("text", "").strip())
    if not counts:
        return 11.0, []
    body = max(counts, key=counts.get)
    heading_sizes = sorted((size for size in counts if size >= body * ratio), reverse=True)
    return body, heading_sizes[:4]


def _heading_level(text, size, bold_ratio, body_size, heading_levels, max_chars) -> int | None:
    if len(text) > max_chars or text.count("\n") > 2 or text.endswith((".", ";", ",")):
        return None
    if re.fullmatch(r"[\d\s.,/-]+", text):
        return None  # Page numbers and numeric labels.
    for index, heading_size in enumerate(heading_levels):
        if size >= heading_size - 0.05:
            return index + 1
    if bold_ratio > 0.8 and size >= body_size - 0.05 and len(text) <= 120:
        numbered = re.match(r"^(\d+(\.\d+)*)[.)]?\s", text)
        if numbered:
            return min(numbered.group(1).count(".") + 1 + len(heading_levels), 6)
        return len(heading_levels) + 1
    return None


def _docx_heading_level(paragraph, style_name: str) -> int | None:
    name = (style_name or "").lower()
    if name == "title":
        return 1
    match = re.match(r"(heading|tiêu đề|đề mục)\s*(\d+)", name)
    if match:
        return int(match.group(2))
    outline = paragraph._p.find(f".//{{{DRAWING_NS['w']}}}outlineLvl")
    if outline is not None:
        value = outline.get(f"{{{DRAWING_NS['w']}}}val")
        if value is not None and value.isdigit() and int(value) < 9:
            return int(value) + 1
    return None


def iter_docx_blocks(document) -> Iterable:
    from docx.document import Document as DocumentType
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    parent = document.element.body if isinstance(document, DocumentType) else document._tc
    for child in parent.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, document)
        elif isinstance(child, CT_Tbl):
            yield Table(child, document)


def format_cell_value(value, number_format: str | None = None) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, dt.datetime):
        return value.date().isoformat() if value.time() == dt.time(0) else value.isoformat(sep=" ")
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, float):
        if number_format and "%" in number_format:
            return f"{value * 100:g}%"
        return str(int(value)) if value.is_integer() else repr(value)
    return clean_text(str(value))


def split_excel_regions_from_grid(grid: list[list[str]]) -> list[tuple[int, int, int, int]]:
    """Detect separate data regions: blank rows split bands, blank columns split bands."""
    regions = []
    row_used = [any(cell for cell in row) for row in grid]
    start = None
    for index, used in enumerate(row_used + [False]):
        if used and start is None:
            start = index
        elif not used and start is not None:
            band = grid[start:index]
            width = max(len(row) for row in band)
            column_used = [any(c < len(row) and row[c] for row in band) for c in range(width)]
            col_start = None
            for col, col_flag in enumerate(column_used + [False]):
                if col_flag and col_start is None:
                    col_start = col
                elif not col_flag and col_start is not None:
                    rows_in_group = [
                        r for r in range(start, index) if any(grid[r][c] for c in range(col_start, min(col, len(grid[r]))))
                    ]
                    if rows_in_group:
                        regions.append((rows_in_group[0] + 1, rows_in_group[-1] + 1, col_start + 1, col))
                    col_start = None
            start = None
    regions.sort()
    return regions


def split_excel_regions(sheet) -> list[tuple[int, int, int, int]]:
    grid = [["" if value is None else str(value) for value in row] for row in sheet.iter_rows(values_only=True)]
    return split_excel_regions_from_grid(grid)


def classify_excel_region(rows: list[list[str]]) -> str:
    cells = [cell for row in rows for cell in row if cell]
    if not cells:
        return "empty"
    numeric = [cell for cell in cells if parse_number(cell, python_repr=True) is not None]
    if len(set(cells)) <= 2 and not numeric and (len(rows) == 1 or len(set(cells)) == 1):
        return "text"  # Titles and notes, including merged title rows.
    width = max(len(row) for row in rows)
    if len(rows) == 1 and width <= 8 and len(cells) >= 2:
        return "kpi" if len(numeric) >= 1 and len(numeric) < len(cells) else "table"
    if len(rows) == 2 and width <= 8:
        header_numeric = sum(1 for cell in rows[0] if cell and parse_number(cell, python_repr=True) is not None)
        value_numeric = sum(1 for cell in rows[1] if cell and parse_number(cell, python_repr=True) is not None)
        if header_numeric == 0 and value_numeric >= max(1, len([c for c in rows[1] if c]) // 2):
            return "kpi"
    if width == 2 and len(rows) <= 4 and all(row[0] and len(row) > 1 and row[1] for row in rows):
        if all(parse_number(row[1], python_repr=True) is not None for row in rows):
            return "kpi"
    return "table"


def render_kpi(rows: list[list[str]]) -> str:
    pairs = []
    if len(rows) == 2:
        pairs = [(label, value) for label, value in zip(rows[0], rows[1]) if label or value]
    elif all(len(row) == 2 for row in rows) or len(rows) == 1 and len(rows[0]) == 2:
        pairs = [(row[0], row[1]) for row in rows]
    else:
        flat = [cell for row in rows for cell in row if cell]
        pairs = list(zip(flat[0::2], flat[1::2]))
    return "KPI:\n" + "\n".join(f"- {label}: {value}" for label, value in pairs)


def openpyxl_chart_info(chart, value_workbook) -> dict:
    info = {
        "chart_type": type(chart).__name__,
        "title": _openpyxl_title(getattr(chart, "title", None)),
        "x_axis": _openpyxl_title(getattr(getattr(chart, "x_axis", None), "title", None)),
        "y_axis": _openpyxl_title(getattr(getattr(chart, "y_axis", None), "title", None)),
        "source_ranges": extract_chart_references(chart),
        "series": [],
    }
    for series in getattr(chart, "ser", []):
        name_ref = _nested(series, ("tx", "strRef", "f"))
        values_ref = _nested(series, ("val", "numRef", "f"))
        category_ref = _nested(series, ("cat", "strRef", "f")) or _nested(series, ("cat", "numRef", "f"))
        name_values = resolve_range_values(value_workbook, name_ref)
        info["series"].append(
            {
                "name": name_values[0] if name_values else (_nested(series, ("tx", "v")) or ""),
                "categories": resolve_range_values(value_workbook, category_ref),
                "values": resolve_range_values(value_workbook, values_ref),
                "values_ref": values_ref,
                "categories_ref": category_ref,
            }
        )
    return info


def _openpyxl_title(title) -> str:
    if title is None:
        return ""
    if isinstance(title, str):
        return title
    texts = []
    rich = getattr(getattr(title, "tx", None), "rich", None)
    for paragraph in getattr(rich, "p", []) or []:
        for run in getattr(paragraph, "r", []) or []:
            if getattr(run, "t", None):
                texts.append(run.t)
    return clean_text("".join(texts))


def _nested(value, path):
    for attribute in path:
        value = getattr(value, attribute, None)
        if value is None:
            return None
    return value


def resolve_range_values(workbook, reference: str | None) -> list[str]:
    parsed = parse_cell_range(reference or "")
    if not parsed:
        return []
    sheet_name, min_row, max_row, min_col, max_col = parsed
    try:
        sheet = workbook[sheet_name] if sheet_name else workbook.active
    except KeyError:
        return []
    values = []
    for row in sheet.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        for cell in row:
            values.append(format_cell_value(cell.value, cell.number_format))
    return values


def extract_chart_references(chart) -> list[str]:
    references = []
    for series in getattr(chart, "ser", []):
        for attribute_path in (
            ("val", "numRef", "f"),
            ("cat", "numRef", "f"),
            ("cat", "strRef", "f"),
            ("tx", "strRef", "f"),
        ):
            value = _nested(series, attribute_path)
            if value:
                references.append(str(value))
    return list(dict.fromkeys(references))


def parse_chart_xml(blob: bytes) -> dict | None:
    """Read chart type, title, axis titles and cached series data from DrawingML chart XML."""
    try:
        root = ElementTree.fromstring(blob)
    except ElementTree.ParseError:
        return None
    ns = {"c": DRAWING_NS["c"], "a": DRAWING_NS["a"]}
    plot_area = root.find(".//c:plotArea", ns)
    if plot_area is None:
        return None

    def rich_text(node) -> str:
        if node is None:
            return ""
        return clean_text("".join(text.text or "" for text in node.iter(f"{{{ns['a']}}}t")))

    chart_types = [
        child.tag.split("}")[1] for child in plot_area if child.tag.endswith("Chart") and "}" in child.tag
    ]
    axes = [rich_text(axis.find("c:title", ns)) for axis in plot_area if axis.tag.endswith(("catAx", "valAx", "dateAx"))]
    info = {
        "chart_type": ", ".join(chart_types) or "chart",
        "title": rich_text(root.find(".//c:chart/c:title", ns)),
        "x_axis": axes[0] if axes else "",
        "y_axis": axes[1] if len(axes) > 1 else "",
        "source_ranges": [],
        "series": [],
    }
    for series in plot_area.iter(f"{{{ns['c']}}}ser"):
        def cache_points(tag):
            node = series.find(f"c:{tag}", ns)
            if node is None:
                return [], None
            formula = node.find(".//c:f", ns)
            points = sorted(
                ((int(pt.get("idx", 0)), (pt.findtext("c:v", default="", namespaces=ns))) for pt in node.iter(f"{{{ns['c']}}}pt")),
                key=lambda item: item[0],
            )
            return [value for _, value in points], formula.text if formula is not None else None

        names, name_ref = cache_points("tx")
        categories, category_ref = cache_points("cat")
        values, values_ref = cache_points("val")
        info["series"].append(
            {
                "name": names[0] if names else "",
                "categories": categories,
                "values": values,
                "values_ref": values_ref,
                "categories_ref": category_ref,
            }
        )
        info["source_ranges"].extend(ref for ref in (values_ref, category_ref, name_ref) if ref)
    info["source_ranges"] = list(dict.fromkeys(info["source_ranges"]))
    return info


def render_chart(info: dict, caption: str = "") -> str:
    lines = []
    if caption:
        lines.append(f"Caption: {caption}")
    if info.get("title"):
        lines.append(f"Chart title: {info['title']}")
    lines.append(f"Chart type: {info.get('chart_type', 'chart')}")
    if info.get("x_axis"):
        lines.append(f"Trục X: {info['x_axis']}")
    if info.get("y_axis"):
        lines.append(f"Trục Y: {info['y_axis']}")
    legend = [series.get("name") for series in info.get("series", []) if series.get("name")]
    if legend:
        lines.append("Chú giải: " + ", ".join(legend))
    data_lines = []
    for series in info.get("series", []):
        values = series.get("values") or []
        categories = series.get("categories") or [str(index + 1) for index in range(len(values))]
        if not values:
            continue
        pairs = ", ".join(f"{category}: {value}" for category, value in zip(categories, values))
        data_lines.append(f"- {series.get('name') or 'Series'}: {pairs}")
    if data_lines:
        lines.append("Dữ liệu nguồn (lấy từ workbook/chart cache):\n" + "\n".join(data_lines[:20]))
    if info.get("source_ranges"):
        lines.append("Source ranges: " + ", ".join(info["source_ranges"]))
    return "\n".join(lines)


def save_parsed_document(document: ParsedDocument, output_dir: Path) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    path = output_dir / f"{document.document_id}.json"
    path.write_text(json.dumps(document.to_dict(), ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    return path


def save_table_parquet(document: ParsedDocument, table_dir: Path) -> int:
    """Persist table rows to Parquet so structured computation does not depend on prompts."""
    try:
        import pandas as pd
    except ImportError:
        return 0
    written = 0
    target_dir = table_dir / document.document_id
    for block in document.blocks:
        rows = block.metadata.get("rows")
        if block.block_type not in ("table", "kpi") or not rows or len(rows) < 2:
            continue
        header = _unique_headers(rows[0])
        width = len(header)
        body = [list(row) + [""] * (width - len(row)) for row in rows[1:]]
        frame = pd.DataFrame([row[:width] for row in body], columns=header)
        target_dir.mkdir(parents=True, exist_ok=True)
        path = target_dir / f"{block.block_id}.parquet"
        try:
            frame.to_parquet(path, index=False)
        except Exception as exc:  # pyarrow missing or unsupported types.
            LOGGER.warning("Could not write parquet for %s: %s", block.block_id, exc)
            return written
        block.metadata["table_path"] = path.relative_to(table_dir.parent).as_posix()
        written += 1
    return written


def _unique_headers(header: list[str]) -> list[str]:
    seen: dict[str, int] = {}
    result = []
    for index, name in enumerate(header):
        name = str(name).strip() or f"column_{index + 1}"
        if name in seen:
            seen[name] += 1
            name = f"{name}_{seen[name]}"
        else:
            seen[name] = 0
        result.append(name)
    return result
